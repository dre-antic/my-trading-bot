from __future__ import annotations

from io import BytesIO
from pathlib import Path


class ExtractedDocument:
    def __init__(self, text: str, filename: str, warning: str | None = None) -> None:
        self.text = text
        self.filename = filename
        self.warning = warning


def extract_document(
    *,
    filename: str = "pasted.txt",
    data: bytes | None = None,
    text: str | None = None,
) -> ExtractedDocument:
    if text and text.strip():
        return ExtractedDocument(text.strip(), filename)

    if data is None:
        raise ValueError("Upload a file or paste your trading style.")

    lower = filename.lower()
    if lower.endswith(".pdf"):
        body, warning = _pdf_text(data)
    elif lower.endswith(".docx"):
        body = _docx_text(data)
        warning = None
    else:
        body = data.decode("utf-8", errors="replace")
        warning = None

    body = body.strip()
    if len(body) < 40:
        raise ValueError(
            "Could not read enough text from that file. If it is a scanned PDF, "
            "paste the strategy in the text box instead."
        )
    return ExtractedDocument(body, filename, warning)


def _pdf_text(data: bytes) -> tuple[str, str | None]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    text = "\n".join(pages).strip()
    warning = None
    if len(text) < 40:
        warning = "This PDF looks like a scan. Paste the text of your style instead."
    return text, warning


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def read_path(path: Path) -> ExtractedDocument:
    return extract_document(filename=path.name, data=path.read_bytes())
